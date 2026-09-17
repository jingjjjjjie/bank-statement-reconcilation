"""Load and apply field-based Excel formatting independently of bank values."""
from copy import copy
from xml.etree import ElementTree as ET

from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from openpyxl.utils import get_column_letter


ROW_NAMES = ("title", "subtitle", "spacer", "header", "opening", "transaction", "total", "sql", "variance")
COMPONENTS = {"font": Font, "fill": PatternFill, "border": Border,
              "alignment": Alignment, "protection": Protection}


class WorkbookStyle:
    """Resolve ordered fields and named row styles before writing any workbook."""

    def __init__(self, path, fields):
        """Reject missing, duplicated, or unknown mappings instead of shifting data."""
        root = ET.parse(path).getroot()
        if root.tag != "workbook-style" or root.get("version") != "2":
            raise ValueError("Unsupported workbook style; expected version 2")
        self.columns = [dict(column.attrib) for column in root.findall("columns/column")]
        names = [column.get("field") for column in self.columns]
        if len(names) != len(fields) or set(names) != set(fields):
            raise ValueError("Style columns must include each supported field exactly once")
        self.positions = {name: index for index, name in enumerate(names, 1)}
        self.last_column = get_column_letter(len(names))
        self.styles = {}
        for entry in root.findall("styles/style"):
            name = entry.attrib["name"]
            if name in self.styles:
                raise ValueError(f"Duplicate cell style: {name}")
            parts = {"number_format": entry.attrib["number_format"]}
            for attribute, factory in COMPONENTS.items():
                component = entry.find(attribute)
                if component is None:
                    raise ValueError(f"Style {name} is missing {attribute}")
                parts[attribute] = factory.from_tree(component)
            self.styles[name] = parts
        self.rows = {}
        self.anchors = {}
        for row in root.findall("row"):
            name = row.attrib["name"]
            entries = row.findall("cell")
            mapping = {entry.get("field"): entry.get("style") for entry in entries}
            if name in self.rows or len(entries) != len(fields) or set(mapping) != set(fields):
                raise ValueError(f"Invalid or duplicate row mapping: {name}")
            if any(style not in self.styles for style in mapping.values()):
                raise ValueError(f"Unknown cell style in row: {name}")
            self.rows[name] = (float(row.attrib["height"]), mapping)
            if name in ("title", "subtitle"):
                anchor = row.get("anchor-style")
                if anchor not in self.styles:
                    raise ValueError(f"Missing or unknown anchor style: {name}")
                self.anchors[name] = anchor
        if set(self.rows) != set(ROW_NAMES):
            raise ValueError("Workbook style must define each named row exactly once")
        self.layout = root.find("print").attrib
        theme = root.find("{http://schemas.openxmlformats.org/drawingml/2006/main}theme")
        self.theme = ET.tostring(theme, encoding="utf-8") if theme is not None else None

    def cell(self, sheet, row, field):
        """Locate a semantic field in the configured column order."""
        return sheet.cell(row, self.positions[field])

    def apply_row(self, sheet, number, name):
        """Keep each field's appearance attached when columns move."""
        height, mapping = self.rows[name]
        sheet.row_dimensions[number].height = height
        for field, style in mapping.items():
            cell = self.cell(sheet, number, field)
            for attribute, value in self.styles[style].items():
                setattr(cell, attribute, copy(value))

    def write_row(self, sheet, number, values):
        """Write typed values by field name, treating strings as literal text."""
        for field, value in values.items():
            cell = self.cell(sheet, number, field)
            cell.value = value
            if isinstance(value, str):
                cell.data_type = "s"

    def prepare(self, book):
        """Apply the workbook theme, widths, and field headings."""
        if self.theme is not None:
            book.loaded_theme = self.theme
        sheet = book.active
        for index, column in enumerate(self.columns, 1):
            sheet.column_dimensions[get_column_letter(index)].width = float(column["width"])
        for number, name in enumerate(ROW_NAMES[:5], 1):
            self.apply_row(sheet, number, name)
        for number, name in ((1, "title"), (2, "subtitle")):
            for attribute, value in self.styles[self.anchors[name]].items():
                setattr(sheet.cell(number, 1), attribute, copy(value))
            sheet.merge_cells(f"A{number}:{self.last_column}{number}")
        self.write_row(sheet, 4, {column["field"]: column["heading"] for column in self.columns})

    def finish(self, sheet, last_row):
        """Apply page layout to the complete dynamically sized report."""
        sheet.page_setup.orientation = self.layout["orientation"]
        sheet.page_setup.paperSize = self.layout["paperSize"]
        sheet.page_setup.fitToWidth = int(self.layout["fitToWidth"])
        sheet.page_setup.fitToHeight = int(self.layout["fitToHeight"])
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.print_title_rows = self.layout["repeat_rows"]
        sheet.print_area = f"A1:{self.last_column}{last_row}"
