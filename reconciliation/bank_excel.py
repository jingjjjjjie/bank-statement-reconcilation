"""Export validated bank transactions using the editable workbook style."""

import argparse
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Protection
from reconciliation.paths import WORKSPACE
from openpyxl.comments import Comment
from reconciliation.bank_statement import read_master


STYLE_PATH = WORKSPACE / "prompts" / "styles" / "bank_statement.xml"


def apply_row_style(style, name, sheet, number):
    """Apply a named row's formatting without introducing cell values."""
    row = style.find(f"row[@name='{name}']")
    if row is None:
        raise ValueError(f"Workbook style is missing row: {name}")
    sheet.row_dimensions[number].height = float(row.attrib["height"])
    components = {"font": Font, "fill": PatternFill, "border": Border,
                  "alignment": Alignment, "protection": Protection}
    for entry in row.findall("cell"):
        cell = sheet.cell(number, int(entry.attrib["column"]))
        definition = style.find(f"styles/style[@name='{entry.attrib["style"]}']")
        cell.number_format = definition.attrib["number_format"]
        for attribute, factory in components.items():
            setattr(cell, attribute, factory.from_tree(definition.find(attribute)))


def export(data_path, company, output_path, style_path=None):
    """Write bank evidence using a fresh workbook and an editable style file."""
    if not company.strip():
        raise ValueError("Enter the company name for the workbook")
    data = read_master(Path(data_path))
    style = ET.parse(Path(style_path) if style_path is not None else STYLE_PATH).getroot()
    if style.tag != "workbook-style" or style.get("version") != "1":
        raise ValueError("Unsupported workbook style")
    book = Workbook()
    theme = style.find("{http://schemas.openxmlformats.org/drawingml/2006/main}theme")
    if theme is not None:
        book.loaded_theme = ET.tostring(theme, encoding="utf-8")
    sheet = book.active
    sheet.title = datetime.fromisoformat(data["transactions"][0]["date"]).strftime("%b'%y").upper()

    # Style files contain formatting only; transaction values come from the master.
    for column in style.findall("columns/column"):
        sheet.column_dimensions[column.attrib["letter"]].width = float(column.attrib["width"])
    for row, name in enumerate(("title", "subtitle", "spacer", "header", "opening"), 1):
        apply_row_style(style, name, sheet, row)
    sheet.merge_cells("A1:K1")
    sheet.merge_cells("A2:K2")
    sheet["A1"] = company.strip()
    sheet["A1"].data_type = "s"
    sheet["A2"] = f"AMBANK - BANK & CASH : A/C {data['account']} ({sheet.title}) - {data['currency']}"
    for column in style.findall("columns/column"):
        sheet[f"{column.attrib['letter']}4"] = column.attrib["heading"]
    sheet["J5"] = float(data["opening_balance"])
    sheet["J5"].comment = Comment("Opening balance from the AmBank statement.", "Source")

    # Fill bank fields; reserve PARTICULAR for the supporting-document branch.
    for number, transaction in enumerate(data["transactions"], 6):
        apply_row_style(style, "transaction", sheet, number)
        sheet.cell(number, 1, datetime.fromisoformat(transaction["date"]))
        sheet.cell(number, 6, transaction["counterparty"].upper() or None)
        sheet.cell(number, 6).data_type = "s"
        sheet.cell(number, 8, float(transaction["money_in"]) or None)
        sheet.cell(number, 9, float(transaction["money_out"]) or None)
        sheet.cell(number, 10, float(transaction["balance"]))
        sheet.cell(number, 11, "PENDING" if transaction["counterparty"] else "REVIEW PAY TO")
        sheet.cell(number, 1).comment = Comment(
            f"{Path(data['source']).name}, page {transaction['page']}.\n"
            f"ID: {transaction['transaction_id']}\n{transaction['narration']}\n"
            "Supporting-document matching pending.", "Source")
        sheet.cell(number, 6).comment = Comment(
            f"Role: {transaction['counterparty_role']}\n"
            f"As printed: {transaction['counterparty_raw']}", "Source")

    # Match the footer layout without inventing SQL reconciliation results.
    footer = len(data["transactions"]) + 7
    for offset, name in enumerate(("total", "sql", "variance")):
        apply_row_style(style, name, sheet, footer + offset)
    sheet.cell(footer, 7, "TOTAL")
    sheet.cell(footer, 8, float(data["total_money_in"]))
    sheet.cell(footer, 9, float(data["total_money_out"]))
    sheet.cell(footer + 1, 7, "AS PER SQL")
    sheet.cell(footer + 2, 7, "VARIANCE")

    # Print layout is part of the editable visual style.
    layout = style.find("print").attrib
    sheet.page_setup.orientation = layout["orientation"]
    sheet.page_setup.paperSize = layout["paperSize"]
    sheet.page_setup.fitToWidth = int(layout["fitToWidth"])
    sheet.page_setup.fitToHeight = int(layout["fitToHeight"])
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = layout["repeat_rows"]
    sheet.print_area = f"A1:K{footer + 2}"
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)
    return output_path


def main():
    """Run the workbook exporter from the command line."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("data", type=Path)
    parser.add_argument("--style", type=Path, default=STYLE_PATH)
    parser.add_argument("--company", required=True)
    parser.add_argument("--output", type=Path, default=Path("bank-output/answer_statement.xlsx"))
    args = parser.parse_args()
    print(export(args.data, args.company, args.output, args.style).resolve())


if __name__ == "__main__":
    main()
