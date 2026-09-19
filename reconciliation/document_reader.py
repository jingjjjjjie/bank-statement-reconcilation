"""Extract every page, cell, text part and embedded image for review."""
import io
import json
import zipfile
from xml.etree import ElementTree as ET

import openpyxl
import pymupdf
from PIL import Image, ImageOps, ImageSequence
from reconciliation.review_settings import DEFAULTS, validate
from reconciliation.pdf_routing import MODES, inspect_page


def extract(path, output, config=None):
    """Extract ordered review units and previews without changing the source."""
    output.mkdir(parents=True, exist_ok=True)
    units = []
    config = validate(config if config is not None else dict(DEFAULTS))

    def add(label, text="", image=None, limitation="", blocked="", pdf_probe=None, whole_text=False):
        # Split long text without dropping rows or pages from coverage.
        chunks = [text] if whole_text or pdf_probe is not None else ([text[i:i + 12000] for i in range(0, len(text), 12000)] or [""])
        for part, chunk in enumerate(chunks, 1):
            item = {"label": label if whole_text else f"{label} / part {part}", "text": chunk, "image": None,
                    "limitation": limitation, "blocked": blocked}
            if pdf_probe is not None:
                item["pdf_probe"] = pdf_probe
            if image is not None and (part == 1 or pdf_probe is not None):
                target = output / f"unit-{len(units) + 1:04d}.png"
                picture = ImageOps.exif_transpose(image).convert("RGB")
                picture.thumbnail((2400, 2400))
                picture.save(target)
                item["image"] = str(target.resolve())
            units.append(item)

    suffix = path.suffix.lower()
    if suffix == ".pdf":
        _extract_pdf(path, config, add)
    elif suffix in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}:
        _extract_images(path, config, add)
    elif suffix == ".xlsx":
        _extract_excel(path, add)
    elif suffix == ".docx":
        _extract_word(path, add)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")

    if suffix in {".xlsx", ".docx"}:
        _extract_embedded(path, config, add)
    if not units:
        raise ValueError("No readable units extracted")
    return units


def _extract_pdf(path, config, add):
    """Read page text and render evidence according to PDF settings."""
    with pymupdf.open(path) as document:
        if document.needs_pass:
            raise ValueError("Encrypted PDF requires a password")
        for number, page in enumerate(document, 1):
            text = page.get_text()
            sparse = sum(ch.isalnum() for ch in text) < 20
            experimental = config["pdf_mode"] in MODES
            needs_picture = experimental or config["pdf_mode"] == "vision" or (config["pdf_mode"] == "auto" and sparse)
            picture, blocked = None, ""
            if needs_picture and config["pictures_enabled"]:
                pixmap = page.get_pixmap(dpi=150, alpha=False)
                picture = Image.open(io.BytesIO(pixmap.tobytes("png")))
            elif sparse:
                blocked = "PDF page has insufficient extractable text; enable PDF fallback/full vision and pictures"
            limitation = "" if picture else "PDF text only: images, handwriting, signatures and visual layout were not inspected"
            if needs_picture and not config["pictures_enabled"]:
                blocked = "PDF vision requested but picture processing is off"
            add(f"page {number}", text, picture, limitation, blocked,
                pdf_probe=inspect_page(page) if experimental else None)


def _extract_images(path, config, add):
    """Keep every image frame or record that pictures are disabled."""
    if not config["pictures_enabled"]:
        add("image", blocked="Picture processing is off; this document has not been reviewed")
        return
    with Image.open(path) as picture:
        for number, frame in enumerate(ImageSequence.Iterator(picture), 1):
            add(f"image {number}", image=frame.copy())


def _extract_excel(path, add):
    """Read formulas, cached values, coordinates, and hidden sheets."""
    formulas = openpyxl.load_workbook(path, read_only=True, data_only=False)
    values = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        for sheet in formulas:
            rows = []
            for row, cached in zip(sheet.iter_rows(), values[sheet.title].iter_rows()):
                cells = []
                for cell, value in zip(row, cached):
                    if cell.value is not None:
                        content = {"value": str(cell.value)}
                        if cell.data_type == "f":
                            content["cached"] = str(value.value)
                        cells.append(f"{cell.coordinate}={json.dumps(content, ensure_ascii=False)}")
                if cells:
                    rows.append(" | ".join(cells))
            # Keep headings, complete cells, and totals together for each worksheet.
            add(f"sheet {sheet.title} ({sheet.sheet_state})", "\n".join(rows), whole_text=True)
    finally:
        formulas.close()
        values.close()


def _extract_word(path, add):
    """Read Office text parts in XML order without reconstructing layout."""
    with zipfile.ZipFile(path) as package:
        for name in sorted(package.namelist()):
            if name.startswith("word/") and name.endswith(".xml"):
                tree = ET.fromstring(package.read(name))
                text = "\n".join(node.text or "" for node in tree.iter()
                                 if node.tag.endswith("}t"))
                if text.strip():
                    add(name, text)


def _extract_embedded(path, config, add):
    """Include Office images and reject objects requiring manual rendering."""
    with zipfile.ZipFile(path) as package:
        for name in sorted(package.namelist()):
            if "/embeddings/" in name or "/charts/" in name and name.endswith(".xml"):
                raise ValueError(f"Embedded object needs manual rendering: {name}")
            if "/media/" in name and not name.endswith("/"):
                if config["pictures_enabled"]:
                    with Image.open(io.BytesIO(package.read(name))) as picture:
                        add(name, image=picture)
                else:
                    add(name, blocked="Embedded picture was not reviewed because picture processing is off")
