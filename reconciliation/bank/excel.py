"""Export validated bank transactions using the editable workbook style."""

import argparse
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.comments import Comment

from reconciliation.bank.statement import read_master
from reconciliation.bank.workbook_style import WorkbookStyle
from reconciliation.core.paths import WORKSPACE

STYLE_PATH = WORKSPACE / "prompts" / "excel" / "bank_statement.xml"


FIELDS = (
    "date",
    "ledger",
    "sql",
    "sales_type",
    "voucher",
    "counterparty",
    "particular",
    "money_in",
    "money_out",
    "balance",
    "status",
)


def transaction_values(transaction):
    """Map validated bank evidence to typed fields; accounting fields stay blank."""
    return {
        "date": datetime.fromisoformat(transaction["date"]),
        "counterparty": transaction["counterparty"].upper() or None,
        "money_in": float(transaction["money_in"]) or None,
        "money_out": float(transaction["money_out"]) or None,
        "balance": float(transaction["balance"]),
        "status": "PENDING" if transaction["counterparty"] else "REVIEW PAY TO",
    }


def statement_workbook(company, metadata, rows, style_path=None):
    """Build the shared statement layout from typed bank or final-review rows."""
    style = WorkbookStyle(Path(style_path) if style_path is not None else STYLE_PATH, FIELDS)
    book = Workbook()
    sheet = book.active
    sheet.title = rows[0]['date'].strftime("%b'%y").upper()
    style.prepare(book)
    sheet['A1'] = company.strip()
    sheet['A1'].data_type = 's'
    sheet['A2'] = f"AMBANK - BANK & CASH : A/C {metadata.get('account', '')} ({sheet.title}) - {metadata['currency']}"
    opening = metadata.get('opening_balance')
    style.write_row(sheet, 5, {'balance': float(opening) if opening not in (None, '') else None})
    style.cell(sheet, 5, 'balance').comment = Comment('Opening balance from the AmBank statement.', 'Source')
    for number, values in enumerate(rows, 6):
        style.apply_row(sheet, number, 'transaction')
        style.write_row(sheet, number, values)
    footer = len(rows) + 7
    for offset, name in enumerate(('total', 'sql', 'variance')):
        style.apply_row(sheet, footer + offset, name)
    style.write_row(
        sheet,
        footer,
        {
            'particular': 'TOTAL',
            'money_in': metadata.get('total_money_in'),
            'money_out': metadata.get('total_money_out'),
        },
    )
    style.write_row(sheet, footer + 1, {'particular': 'AS PER SQL'})
    style.write_row(sheet, footer + 2, {'particular': 'VARIANCE'})
    style.finish(sheet, footer + 2)
    return book, style


def export(data_path, company, output_path, style_path=None):
    """Write validated bank evidence using the shared statement layout."""
    if not company.strip():
        raise ValueError('Enter the company name for the workbook')
    data = read_master(Path(data_path))
    metadata = {**data, **{key: float(data[key]) for key in ('total_money_in', 'total_money_out')}}
    book, style = statement_workbook(
        company, metadata, [transaction_values(t) for t in data['transactions']], style_path
    )
    sheet = book.active
    for number, transaction in enumerate(data['transactions'], 6):
        style.cell(sheet, number, 'date').comment = Comment(
            f"{Path(data['source']).name}, page {transaction['page']}.\n"
            f"ID: {transaction['transaction_id']}\n{transaction['narration']}\n"
            'Supporting-document matching pending.',
            'Source',
        )
        style.cell(sheet, number, 'counterparty').comment = Comment(
            f"Role: {transaction['counterparty_role']}\nAs printed: {transaction['counterparty_raw']}", 'Source'
        )
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(output_path)
    return output_path


def main():
    """Run the workbook exporter from the command line."""
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("data", type=Path)
    parser.add_argument("--style", type=Path, default=STYLE_PATH)
    parser.add_argument("--company", required=True)
    parser.add_argument("--output", type=Path, default=Path("bank-output/answer_statement.xlsx"))
    args = parser.parse_args()
    print(export(args.data, args.company, args.output, args.style).resolve())


if __name__ == "__main__":
    main()
