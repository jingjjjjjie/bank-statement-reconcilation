"""Export extracted transactions using the sample workbook's December styles."""

import argparse
from copy import copy
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.comments import Comment
from bank_statement import read_master


def copy_style(source, target):
    """Register each style component in the destination workbook."""
    for attribute in ("font", "fill", "border", "alignment", "protection", "number_format"):
        setattr(target, attribute, copy(getattr(source, attribute)))


def main():
    # Load only extraction data and the chosen style reference.
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("data", type=Path)
    parser.add_argument("template", type=Path)
    parser.add_argument("--company", required=True)
    parser.add_argument("--output", type=Path, default=Path("bank-output/answer_statement.xlsx"))
    args = parser.parse_args()
    data = read_master(args.data)
    source = load_workbook(args.template)
    reference = source["DEC'25"]
    book = Workbook()
    book.loaded_theme = source.loaded_theme
    sheet = book.active
    sheet.title = datetime.fromisoformat(data["transactions"][0]["date"]).strftime("%b'%y").upper()

    # Copy styles into a fresh workbook so no other customer's records survive.
    for key, dimension in reference.column_dimensions.items():
        if key in "ABCDEFGHIJK":
            sheet.column_dimensions[key].width = dimension.width
    for row in range(1, 6):
        for column in range(1, 12):
            copy_style(reference.cell(row, column), sheet.cell(row, column))
        sheet.row_dimensions[row].height = reference.row_dimensions[row].height
    sheet.merge_cells("A1:K1")
    sheet.merge_cells("A2:K2")
    sheet["A1"] = args.company
    sheet["A2"] = f"AMBANK - BANK & CASH : A/C {data['account']} ({sheet.title}) - {data['currency']}"
    for cell in reference[4][:11]:
        sheet.cell(4, cell.column, cell.value)
    sheet["J5"] = float(data["opening_balance"])
    sheet["J5"].comment = Comment("Opening balance from the AmBank statement.", "Source")

    # Fill bank fields; reserve PARTICULAR for the supporting-document branch.
    for number, transaction in enumerate(data["transactions"], 6):
        for column in range(1, 12):
            copy_style(reference.cell(6, column), sheet.cell(number, column))
        sheet.cell(number, 1, datetime.fromisoformat(transaction["date"]))
        sheet.cell(number, 1).number_format = "dd/mm/yyyy"
        sheet.cell(number, 6, transaction["counterparty"].upper() or None)
        sheet.cell(number, 6).data_type = "s"
        sheet.cell(number, 8, float(transaction["money_in"]) or None)
        sheet.cell(number, 9, float(transaction["money_out"]) or None)
        sheet.cell(number, 10, float(transaction["balance"]))
        sheet.cell(number, 11, "PENDING" if transaction["counterparty"] else "REVIEW PAY TO")
        sheet.cell(number, 1).comment = Comment(
            f"{Path(data['source']).name}, page {transaction['page']}.\n"
            f"ID: {transaction['transaction_id']}\n{transaction['narration']}\n"
            "Supporting-document matching pending.", "Source")
        sheet.cell(number, 6).comment = Comment(
            f"Role: {transaction['counterparty_role']}\n"
            f"As printed: {transaction['counterparty_raw']}", "Source")
        sheet.row_dimensions[number].height = 30

    # Match the footer layout without inventing SQL reconciliation results.
    footer = len(data["transactions"]) + 7
    for offset, source_row in enumerate((160, 161, 162)):
        for column in range(1, 12):
            copy_style(reference.cell(source_row, column), sheet.cell(footer + offset, column))
    sheet.cell(footer, 7, "TOTAL")
    sheet.cell(footer, 8, float(data["total_money_in"]))
    sheet.cell(footer, 9, float(data["total_money_out"]))
    sheet.cell(footer + 1, 7, "AS PER SQL")
    sheet.cell(footer + 2, 7, "VARIANCE")

    # Fit the wide report to paper while retaining the sample's column widths.
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "4:4"
    sheet.print_area = f"A1:K{footer + 2}"
    args.output.parent.mkdir(parents=True, exist_ok=True)
    book.save(args.output)
    source.close()
    print(args.output.resolve())


if __name__ == "__main__":
    main()
