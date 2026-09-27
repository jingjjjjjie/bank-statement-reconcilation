"""Verify source-row navigation against an actual workbook preview."""
import json
import threading
import unittest
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import PatternFill
from playwright.sync_api import expect, sync_playwright

from dashboard.routes import create_app
from reconciliation.duplicate_workflow import fingerprint
from tests.browser import browser_options
from tests.http_server import TestServer
from tests.unit import test_receipt_matching as fixtures


class ExcelHighlightTests(unittest.TestCase):
    def test_show_source_rows_without_editing(self):
        """Jump across worksheet chunks, highlight exact rows and clear stale markers."""
        fixture = fixtures.ReceiptMatchingTests()
        fixture.setUp()
        self.addCleanup(fixture.doCleanups)
        source = fixture.base / 'payments.xlsx'
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Sheet A'
        for row in range(1, 101):
            sheet.append([f'Person {row}', f'Reference {row}', row])
        sheet['C55'].fill = PatternFill('solid', fgColor='00FF00')
        workbook.create_sheet('Other').append(['Header', 'Amount'])
        workbook['Other'].append(['Other person', 15])
        workbook.save(source)
        digest = fingerprint(source)
        fixture.index['documents'] = {digest: {'paths': [str(source)], 'error': None,
            'units': [{'label': 'sheet Sheet A (visible)', 'image': None}]}}
        (fixture.work / 'index.json').write_text(json.dumps(fixture.index))
        pieces = [{**fixture.pieces[0], 'amount_location': "'Sheet A'!$A$55:$C$56"},
                  {**fixture.pieces[1], 'amount_location': 'Other!B2'}]
        fixture.state['units'] = {digest + ':0': {'readable': True, 'receipts': pieces}}
        fixture.state['index_sha256'] = fingerprint(fixture.work / 'index.json')
        fixture.save_state()
        review = fixture.review
        review.root, review.manifest, review.data = fixture.base, {}, fixture.base / 'data'
        review.data.mkdir()
        review.workspace = lambda: {'name': 'Fixture', 'period': ''}
        review.workflow_checks = lambda: (False, False, False)
        server = TestServer(('127.0.0.1', 0), create_app(review, 'token'))
        threading.Thread(target=server.serve_forever, daemon=True).start()
        self.addCleanup(server.server_close)
        self.addCleanup(server.shutdown)
        with sync_playwright() as playwright:
            browser = playwright.chromium.launch(**browser_options())
            page = browser.new_page(viewport={'width': 1440, 'height': 960})
            page.goto(f'http://127.0.0.1:{server.server_port}/extraction-review')
            expect(page.locator('.sheet-table')).to_be_visible()
            page.get_by_role('button', name='Show original for piece 1', exact=True).click()
            expect(page.locator('#original-page')).to_have_value('1')
            expect(page.locator('.source-highlight')).to_have_count(2)
            expect(page.locator('.source-highlight').first).to_have_attribute('data-source-row', '55')
            expect(page.locator('.source-highlight').last).to_have_attribute('data-source-row', '56')
            expect(page.locator('.source-highlight td').first).to_have_css('background-color', 'rgb(255, 240, 174)')
            expect(page.locator('#document-review-status')).to_have_attribute('aria-label', 'Not reviewed')
            viewport = page.locator('#original-viewport').bounding_box()
            row = page.locator('.source-highlight').first.bounding_box()
            self.assertGreaterEqual(row['y'], viewport['y'])
            self.assertLessEqual(row['y'] + row['height'], viewport['y'] + viewport['height'])
            folder = Path('duplicated/inspection')
            folder.mkdir(parents=True, exist_ok=True)
            page.screenshot(path=str(folder / 'excel-row-highlight.png'))
            page.set_viewport_size({'width':390, 'height':844})
            page.get_by_role('button', name='Show original for piece 1', exact=True).click()
            expect(page.locator('.source-highlight')).to_have_count(2)
            self.assertTrue(page.evaluate('document.documentElement.scrollWidth <= innerWidth'))
            page.set_viewport_size({'width':1440, 'height':960})
            page.get_by_role('button', name='Show original for piece 2', exact=True).click()
            expect(page.locator('.sheet-title')).to_have_text('Other')
            expect(page.locator('.source-highlight')).to_have_count(1)
            expect(page.locator('.source-highlight')).to_have_attribute('data-source-row', '2')
            page.locator('#original-page').select_option('0')
            expect(page.locator('.sheet-title')).to_have_text('Sheet A')
            expect(page.locator('.source-highlight')).to_have_count(0)
            page.get_by_role('button', name='Show original for piece 1', exact=True).click()
            expect(page.locator('.source-highlight')).to_have_count(2)
            page.locator('.piece-details summary').first.click()
            page.locator('[data-field="amount_location"]').first.fill('Missing!B55')
            page.get_by_role('button', name='Show original for piece 1', exact=True).click()
            expect(page.locator('#original-status')).to_contain_text('Location not found')
            expect(page.locator('.source-highlight')).to_have_count(0)
            browser.close()
