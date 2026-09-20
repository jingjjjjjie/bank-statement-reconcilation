"""Verify print preview fidelity, reuse, invalidation, and display-only conversion."""
import os
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

import pymupdf
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from dashboard import excel_pdf, extraction_preview


@unittest.skipUnless(shutil.which("libreoffice") or shutil.which("soffice"), "LibreOffice required")
class ExcelPdfTests(unittest.TestCase):
    def setUp(self):
        """Build a styled two-sheet workbook with explicit print areas."""
        folder = tempfile.TemporaryDirectory()
        self.addCleanup(folder.cleanup)
        self.base = Path(folder.name)
        self.addCleanup(patch.stopall)
        patch.dict(os.environ, {"EXCEL_PREVIEW_CACHE": str(self.base / "cache")}).start()
        self.source = self.base / "original.xlsx"
        self.book = Workbook()
        sheet = self.book.active
        sheet.title = "Invoice"
        sheet.merge_cells("A1:D1")
        sheet["A1"] = "INVOICE PREVIEW"
        sheet["A1"].font = Font(name="Carlito", size=20, bold=True, color="FFFFFF")
        sheet["A1"].fill = PatternFill("solid", fgColor="234567")
        sheet.row_dimensions[1].height = 36
        sheet["A3"] = "Total"
        sheet["C3"] = 1234.5
        sheet["C3"].number_format = '#,##0.00'
        sheet.print_area = "A1:D6"
        sheet["A20"] = "OUTSIDE PRINT AREA"
        other = self.book.create_sheet("Second sheet")
        other["A1"] = "SECOND SHEET"
        other.print_area = "A1:D6"
        self.book.save(self.source)

    def test_cached_print_pages_preserve_original_and_invalidate_on_edit(self):
        """Use real print pages, reuse conversion, and never write into the workbook."""
        original = self.source.read_bytes()
        output = excel_pdf.convert_excel_to_pdf(self.source)
        with pymupdf.open(output) as pdf:
            self.assertEqual(len(pdf), 2)
            text = "".join(page.get_text() for page in pdf)
        self.assertIn("INVOICE PREVIEW", text)
        self.assertIn("1,234.50", text)
        self.assertIn("SECOND SHEET", text)
        self.assertNotIn("OUTSIDE PRINT AREA", text)
        with patch.object(excel_pdf.subprocess, "run", side_effect=AssertionError("cache miss")):
            self.assertEqual(excel_pdf.convert_excel_to_pdf(self.source), output)
            info = extraction_preview.describe(self.source)
            self.assertEqual((info["kind"], info["pages"]), ("pdf", 2))
            self.assertTrue(extraction_preview.image(self.source, 1).startswith(b"\x89PNG"))
        self.assertEqual(self.source.read_bytes(), original)
        self.book.active["A3"] = "Changed total"
        self.book.save(self.source)
        self.assertNotEqual(excel_pdf.convert_excel_to_pdf(self.source), output)

    def test_failed_conversion_is_not_cached(self):
        """Failed conversions remain explicit and can be retried successfully."""
        excel_pdf.converter()
        with patch.object(excel_pdf.subprocess, "run") as run:
            run.return_value.returncode = 1
            with self.assertRaisesRegex(ValueError, "conversion failed"):
                excel_pdf.convert_excel_to_pdf(self.source)
        self.assertEqual(list((self.base / "cache").glob("*.pdf")), [])
        self.assertTrue(excel_pdf.convert_excel_to_pdf(self.source).is_file())
