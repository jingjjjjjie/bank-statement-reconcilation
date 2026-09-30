"""Check styled final exports against the shared bank workbook and saved ledger."""

import csv
import io
from unittest.mock import patch

from openpyxl import load_workbook

from dashboard.services.matching import final_export, final_review
from reconciliation.bank import excel
from tests.fixtures.final_review import FinalReviewFixture


class FinalExportTests(FinalReviewFixture):
    """Use synthetic bound bank amounts and approvals without model calls."""

    def setUp(self):
        """Add statement metadata to the frozen fixture before creating any decisions."""
        super().setUp()
        master = self.project / 'bank-output/master_statement.csv'
        with master.open(encoding='utf-8', newline='') as stream:
            rows = list(csv.DictReader(stream))
        for row, balance in zip(rows, (90, 80, 75)):
            row.update(
                account='123',
                opening_balance='100',
                balance=str(balance),
                money_in='0',
                money_out='5' if balance == 75 else '10',
                transaction_id='tx-' + row['sequence'],
            )
        with master.open('w', encoding='utf-8', newline='') as stream:
            writer = csv.DictWriter(stream, fieldnames=list(rows[0]))
            writer.writeheader()
            writer.writerows(rows)
        facts = final_review.read(self.cache / 'facts.json')
        facts['bank_hash'] = final_review.source_hash(master)
        self.write(self.cache / 'facts.json', facts)
        self.review.export_defaults = lambda: {'company': 'Example company'}

    def workbook(self, paths=True):
        """Read the actual XLSX bytes produced by the final exporter."""
        book = load_workbook(io.BytesIO(final_export.export_workbook(self.review, paths)))
        self.addCleanup(book.close)
        return book

    def test_same_bank_layout_and_optional_paths(self):
        """The statement keeps its format, amounts and totals; paths belong on the details sheet."""
        final_review.decide(self.review, self.request(note='=literal note'))
        banks = final_review.snapshot(self.review)['banks']
        data = {
            'account': '123',
            'currency': 'MYR',
            'source': 'statement.txt',
            'opening_balance': '100',
            'total_money_in': '0',
            'total_money_out': '25',
            'transactions': [
                {
                    **bank,
                    'counterparty': bank['parties'][0],
                    'narration': bank['description'],
                    'counterparty_role': 'recipient',
                    'counterparty_raw': bank['parties'][0],
                }
                for bank in banks
            ],
        }
        baseline_path = self.root / 'baseline.xlsx'
        with patch.object(excel, 'read_master', return_value=data):
            excel.export(self.root / 'unused.csv', 'Example company', baseline_path)
        baseline = load_workbook(baseline_path)
        self.addCleanup(baseline.close)
        for paths in (True, False):
            book = self.workbook(paths)
            sheet = book.active
            self.assertEqual(sheet.max_column, 11)
            self.assertEqual(sheet.merged_cells, baseline.active.merged_cells)
            self.assertEqual(sheet.print_area, baseline.active.print_area)
            self.assertEqual(sheet.page_setup, baseline.active.page_setup)
            for row in baseline.active:
                for cell in row:
                    actual = sheet[cell.coordinate]
                    self.assertEqual(actual._style, cell._style)
                    if not (cell.row in (6, 7, 8) and cell.column in (7, 11)):
                        self.assertEqual(actual.value, cell.value)
            self.assertEqual([sheet.cell(row, 11).value for row in (6, 7, 8)], ['OK', None, None])
            self.assertEqual(sheet['G6'].value, 'Receipt')
            self.assertEqual(sheet['J5'].value, 100)
            self.assertEqual(sheet['I10'].value, 25)
            details = book['Review details']
            self.assertEqual(details.max_row, 4)
            self.assertEqual(details.max_column, 13 if paths else 12)
            self.assertEqual(details['L2'].value, '=literal note')
            self.assertEqual(details['L2'].data_type, 's')
            if paths:
                self.assertEqual(details['M2'].value, str(final_review.evidence(self.review, 'item', 'D1')))
                self.assertIsNone(details['M3'].value)
            self.assertFalse(any(cell.data_type == 'f' for ws in book for row in ws for cell in row))

    def test_partial_and_stale_approvals_keep_blank_remarks(self):
        """Incomplete or changed support never receives OK and differences remain visible."""
        final_review.decide(self.review, self.request(allocations=[{'item_id': 'D1', 'amount': '4'}]))
        book = self.workbook()
        self.assertIsNone(book.active['K6'].value)
        self.assertEqual(book['Review details']['H2'].value, 6)
        final_review.decide(self.review, self.request())
        final_review.evidence(self.review, 'item', 'D1').write_text('Changed original', encoding='utf-8')
        book = self.workbook()
        self.assertIsNone(book.active['K6'].value)
        self.assertIsNone(book.active['G6'].value)
        self.assertTrue(book['Review details']['I2'].value)

    def test_old_import_recovers_only_hash_bound_layout_fields(self):
        """Historical imports can recover balances without accepting changed master bytes."""
        banks = final_review.snapshot(self.review)['banks']
        imported = [
            {key: value for key, value in bank.items() if key not in ('balance', 'opening_balance', 'account')}
            for bank in banks
        ]
        restored = final_export.statement_banks(imported)
        self.assertEqual(restored[0]['balance'], '90')
        self.assertEqual(restored[0]['opening_balance'], '100')
        self.assertNotIn('balance', imported[0])
        master = self.project / 'bank-output/master_statement.csv'
        master.write_text(master.read_text() + '\n', encoding='utf-8')
        self.assertEqual(final_export.statement_banks(imported), imported)
