"""Check reusable styles and bank-only export content without a sample workbook."""
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from reconciliation.bank_excel import STYLE_PATH, export


class BankExcelTests(unittest.TestCase):
    """Exercise saved workbook values and live style changes."""

    def setUp(self):
        """Supply synthetic validated bank data independently of PDF extraction."""
        temporary = tempfile.TemporaryDirectory()
        self.addCleanup(temporary.cleanup)
        self.root = Path(temporary.name)
        self.output = self.root / "answer.xlsx"
        row = {"date": "2026-01-02", "counterparty": "Example vendor", "money_in": "0",
               "money_out": "20", "balance": "80", "page": 2, "transaction_id": "test-1",
               "narration": "Synthetic transfer", "counterparty_role": "recipient",
               "counterparty_raw": "Example vendor"}
        data = {"account": "123", "currency": "MYR", "source": "statement.pdf",
                "opening_balance": "100", "total_money_in": "0", "total_money_out": "20",
                "transactions": [row]}
        mock = patch("reconciliation.bank_excel.read_master", return_value=data)
        mock.start()
        self.addCleanup(mock.stop)

    def test_export_preserves_bank_fields_and_captured_style(self):
        """The saved file contains evidence and formatting without sample data or formulas."""
        export(self.root / "master.csv", "Example company", self.output)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        sheet = book.active
        self.assertEqual(sheet.title, "JAN'26")
        self.assertEqual(sheet["A1"].value, "Example company")
        self.assertEqual(sheet["F6"].value, "EXAMPLE VENDOR")
        self.assertEqual(sheet["I6"].value, 20)
        self.assertEqual(sheet["J6"].value, 80)
        self.assertEqual(sheet["J5"].value, 100)
        self.assertEqual(sheet["K6"].value, "PENDING")
        self.assertEqual(sheet["I8"].value, 20)
        for column in (2, 3, 4, 5, 7):
            self.assertIsNone(sheet.cell(6, column).value)
        self.assertIsNone(sheet["H9"].value)
        self.assertIsNone(sheet["H10"].value)
        self.assertFalse(any(cell.data_type == "f" for row in sheet for cell in row))
        self.assertIn("test-1", sheet["A6"].comment.text)
        self.assertEqual(sheet["A6"].number_format, "dd/mm/yyyy")
        self.assertEqual(sheet["A1"].font.sz, 16)
        self.assertEqual(sheet.row_dimensions[6].height, 30)
        self.assertEqual(sheet.page_setup.orientation, "landscape")
        self.assertIsNotNone(book.loaded_theme)

    def test_style_edits_apply_to_next_export(self):
        """Changing the style file changes formatting without changing bank values."""
        path = self.root / "style.xml"
        tree = ET.parse(STYLE_PATH)
        style = tree.getroot()
        style.find("columns/column[@letter='F']").set("width", "42")
        style.find("row[@name='transaction']").set("height", "44")
        tree.write(path)
        export(self.root / "master.csv", "Example company", self.output, path)
        book = load_workbook(self.output)
        self.addCleanup(book.close)
        self.assertEqual(book.active.column_dimensions["F"].width, 42)
        self.assertEqual(book.active.row_dimensions[6].height, 44)
        self.assertEqual(book.active["J6"].value, 80)

    def test_missing_or_invalid_style_fails_export(self):
        """No hidden sample or fallback format masks a broken style definition."""
        path = self.root / "missing.xml"
        with self.assertRaises(FileNotFoundError):
            export(self.root / "master.csv", "Example company", self.output, path)
        path.write_text('<workbook-style version="99"/>')
        with self.assertRaisesRegex(ValueError, "Unsupported workbook style"):
            export(self.root / "master.csv", "Example company", self.output, path)
        self.assertFalse(self.output.exists())
