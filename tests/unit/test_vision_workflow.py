"""Temporary fixtures verify coverage, admin gates, and failure handling."""
import csv
import copy
import hashlib
import json
import tempfile
import threading
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from PIL import Image
from openpyxl import Workbook
import pymupdf

import reconciliation.vision_workflow as workflow
from reconciliation.codex_reviewer import BudgetReached, CodexReviewer, EXTRACTION, ReviewCancelled, SCREEN
from reconciliation.document_reader import extract
from reconciliation.prompts import load_prompt
from reconciliation.duplicate_workflow import organize


class FakeReviewer:
    model = "fixture"

    def ask(self, prompt, schema, images=()):
        # Deterministic results test orchestration, not model accuracy.
        if schema == EXTRACTION:
            return {"receipts": [], "readable": True, "supporting_evidence_status": "potential_support",
                    "supporting_evidence_reason": "Visible transaction details", "document_type": "receipt", "receipt_status": "receipt", "invoice_numbers": ["TEST-1"],
                    "company": [], "brief_description": "fixture",
                    "references": ["TEST-1"],
                    "parties": [], "dates": [], "amounts_and_currencies": ["MYR 1"],
                    "money": [{"amount": "1.00", "currency": "MYR", "role": "grand_total"}],
                    "details": "fixture", "annotations_and_signatures": "none", "limitations": []}
        if schema == SCREEN:
            payload = json.loads(prompt.split("\n", 1)[1])
            return {"comparisons": [{"right_id": key, "candidate": True, "reason": "same reference"}
                                    for key in payload["right"]]}
        return {"classification": "same_document", "confidence": "high",
                "evidence": ["same test reference"], "differences": [], "limitations": []}


class WorkflowTests(unittest.TestCase):
    def test_checkpoint_retries_a_brief_windows_lock(self):
        """A temporary sharing lock does not lose a saved review checkpoint."""
        with tempfile.TemporaryDirectory() as temporary:
            path = Path(temporary) / "state.json"
            original = Path.replace
            attempts = [0]

            def briefly_locked(source, target):
                """Deny the first replacement and allow the retry."""
                attempts[0] += 1
                if attempts[0] == 1:
                    raise PermissionError("file is in use")
                return original(source, target)

            with patch.object(Path, "replace", briefly_locked), patch("reconciliation.vision_workflow.sleep"):
                workflow.save(path, {"units": 1})
            self.assertEqual(workflow.read(path), {"units": 1})
            self.assertEqual(attempts[0], 2)

    def setUp(self):
        # Fixtures live outside the customer folder and clean themselves up.
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / "sources"
        self.root.mkdir()
        self.work = self.base / "review"
        self.manifest = self.base / "manifest.json"
        for name, color in (("a.png", "white"), ("b.png", "gray")):
            Image.new("RGB", (80, 80), color).save(self.root / name)

    def prepared(self, duplicate=False):
        if duplicate:
            (self.root / "copy.png").write_bytes((self.root / "a.png").read_bytes())
        organize(self.root, self.manifest)
        workflow.prepare(self.manifest, self.work)
        return workflow.load(self.work)

    def completed(self):
        index, state = self.prepared()
        workflow.run(self.work, index, state, FakeReviewer())
        return index, state, next(iter(state["pairs"]))

    def test_admin_required_then_keep_both_passes(self):
        index, state, pair = self.completed()
        self.assertTrue(any("admin verdict" in p for p in workflow.gate(index, state)))
        workflow.decide(self.work, index, state, pair, "keep_both", "Admin", "Separate evidence")
        self.assertEqual(workflow.gate(index, state), [])
        workflow.undo_decision(self.work, index, state, pair, "Admin", "Need another look")
        self.assertTrue(any("admin verdict" in p for p in workflow.gate(index, state)))
        self.assertNotIn(pair, workflow.load(self.work)[1]["decisions"])

    def test_inventory_lists_nested_files_and_review_status(self):
        """Every nested source stays visible through pending and reviewed states."""
        nested = self.root / "deep" / "deeper"
        nested.mkdir(parents=True)
        (self.root / "b.png").rename(nested / "b.png")
        index, state = self.prepared()
        path = self.work / "supporting-inventory.csv"

        def rows():
            with path.open(encoding="utf-8-sig", newline="") as stream:
                return list(csv.DictReader(stream))

        self.assertEqual(len(rows()), 2)
        self.assertTrue(any("deep\\deeper" in row["source_path"] or
                            "deep/deeper" in row["source_path"] for row in rows()))
        self.assertEqual({row["status"] for row in rows()}, {"extraction_pending"})
        workflow.run(self.work, index, state, FakeReviewer())
        self.assertEqual({row["status"] for row in rows()}, {"review_pending"})
        self.assertTrue(all(row["duplicate_with"] != "[]" for row in rows()))
        pair = next(iter(state["pairs"]))
        workflow.decide(self.work, index, state, pair, "keep_both", "Admin", "Different evidence")
        self.assertEqual({row["status"] for row in rows()}, {"reviewed_keep_both"})
        self.assertTrue(all(row["duplicate_with"] == "[]" for row in rows()))
        self.assertEqual({row["invoice_numbers"] for row in rows()}, {'["TEST-1"]'})
        self.assertEqual({row["combined_total"] for row in rows()}, {'{"MYR": "1.00"}'})

    def test_unsupported_and_exact_copies_remain_listed(self):
        """List every original while reviewing only accepted formats."""
        nested = self.root / "deep" / "deeper"
        nested.mkdir(parents=True)
        (nested / "notes.txt").write_text("not a receipt", encoding="utf-8")
        (nested / "copy.png").write_bytes((self.root / "a.png").read_bytes())
        index, state = self.prepared()
        with (self.work / "supporting-inventory.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(len(rows), 4)
        self.assertEqual(len([row for row in rows if row["original_path"].endswith(".png")]), 3)
        self.assertEqual(len([row for row in rows if row["exact_duplicate_with"] != "[]"]), 2)
        unsupported = next(row for row in rows if row["original_path"].endswith("notes.txt"))
        self.assertEqual(unsupported["status"], "not_accepted")
        self.assertEqual(unsupported["format_status"], "not_accepted")
        self.assertEqual(unsupported["receipt_status"], "")
        self.assertFalse(any("notes.txt" in doc["paths"][0] and doc["units"] for doc in index["documents"].values()))

    def test_unsupported_format_does_not_block_review(self):
        """An inventoried unsupported file requires no model call or verdict."""
        (self.root / "notes.txt").write_text("notes", encoding="utf-8")
        index, state = self.prepared()
        workflow.run(self.work, index, state, FakeReviewer())
        pair = next(iter(state["pairs"]))
        workflow.decide(self.work, index, state, pair, "keep_both", "Admin", "Different evidence")
        self.assertEqual(workflow.gate(index, state), [])
        with (self.work / "supporting-inventory.csv").open(encoding="utf-8-sig", newline="") as stream:
            rows = list(csv.DictReader(stream))
        self.assertEqual(next(row for row in rows if row["file_format"] == ".txt")["status"], "not_accepted")
        self.assertEqual({row["receipt_status"] for row in rows if row["file_format"] == ".png"}, {"receipt"})

    def test_strong_fields_skip_model_screen_but_compare_originals(self):
        """A local trigger saves screening tokens without granting a verdict."""
        index, state = self.prepared()
        calls = []

        class RecordingReviewer(FakeReviewer):
            def ask(self, prompt, schema, images=()):
                calls.append(schema)
                return super().ask(prompt, schema, images)

        workflow.run(self.work, index, state, RecordingReviewer())
        self.assertNotIn(SCREEN, calls)
        self.assertEqual(len(state["pairs"]), 1)
        self.assertTrue(any("Combined total" in row["reason"] for row in state["screens"].values()))
        self.assertTrue(workflow.gate(index, state))

    def test_duplicate_cleanup_requires_survivor(self):
        index, state, pair = self.completed()
        workflow.decide(self.work, index, state, pair, "keep_left", "Admin", "Confirmed same document")
        self.assertTrue(any("cleanup pending" in p for p in workflow.gate(index, state)))
        left, right = pair.split(":")
        Path(index["documents"][right]["paths"][0]).unlink()
        self.assertEqual(workflow.gate(index, state), [])
        with self.assertRaisesRegex(ValueError, "unapproved missing"):
            workflow.undo_decision(self.work, index, state, pair, "Admin", "Need another look")
        self.assertIn(pair, workflow.load(self.work)[1]["decisions"])
        _, state = workflow.load(self.work)
        Path(index["documents"][left]["paths"][0]).unlink()
        with self.assertRaises(ValueError):
            workflow.gate(index, state)

    def test_pass_one_blocks_model_calls(self):
        index, state = self.prepared(duplicate=True)
        self.assertTrue((self.base / "duplicated").is_dir())
        self.assertFalse((self.root / "duplicated").exists())
        self.assertEqual(len(index["documents"]), 2)
        with self.assertRaisesRegex(ValueError, "Pass-one"):
            workflow.run(self.work, index, state, FakeReviewer())
        self.assertEqual(state["units"], {})

    def test_budget_keeps_review_incomplete(self):
        index, state = self.prepared()
        class LimitedReviewer(FakeReviewer):
            def ask(self, *args, **kwargs):
                raise BudgetReached("test limit")
        with self.assertRaises(BudgetReached):
            workflow.run(self.work, index, state, LimitedReviewer())
        self.assertTrue(workflow.gate(index, state))
        self.assertTrue((self.work / "report.md").exists())

    def test_refresh_keeps_validated_model_cache(self):
        """Repeated testing reuses an identical completed Codex response."""
        self.prepared()
        prompt, model = "Cached fixture extraction", "fixture"
        key = hashlib.sha256(json.dumps([load_prompt("styles") + "\n\n" + prompt, EXTRACTION, model, "default"],
                                        sort_keys=True).encode()).hexdigest()
        folder = self.work / "model-cache" / key
        folder.mkdir(parents=True)
        expected = FakeReviewer().ask(prompt, EXTRACTION)
        (folder / "result.json").write_text(json.dumps(expected), encoding="utf-8")
        workflow.prepare(self.manifest, self.work, refresh=True)
        reviewer = CodexReviewer(self.work, executable="missing-codex", model=model, max_calls=0)
        self.assertEqual(reviewer.ask(prompt, EXTRACTION), expected)
        self.assertEqual(reviewer.calls, 0)

    def test_cancel_terminates_active_codex_call(self):
        """Stopping an active request leaves no valid cached result."""
        started = threading.Event()

        reviewer = CodexReviewer(self.work, executable="fixture-codex", model="fixture", max_calls=1)
        failures = []

        def read():
            try:
                reviewer.ask("Stop this fixture", EXTRACTION)
            except Exception as error:
                failures.append(error)

        def wait_for_stop(command, **kwargs):
            """Block a mocked exec until the real review cancellation flag is set."""
            if command[1:3] == ["login", "status"]:
                return SimpleNamespace(returncode=0, stdout="chatgpt", stderr="")
            started.set()
            reviewer.processes.cancelled.wait(timeout=5)
            raise ReviewCancelled("Review stopped by user")

        with patch.object(reviewer.processes, "run", side_effect=wait_for_stop):
            worker = threading.Thread(target=read)
            worker.start()
            self.assertTrue(started.wait(timeout=5))
            reviewer.cancel()
            worker.join(timeout=5)
        self.assertFalse(worker.is_alive())
        self.assertIsInstance(failures[0], ReviewCancelled)
        self.assertFalse(reviewer.last_result.exists())

    def test_parallel_units_complete_and_checkpoint(self):
        """Read independent units together and save both validated results."""
        index, state = self.prepared()
        barrier = threading.Barrier(2)

        class ParallelReviewer(FakeReviewer):
            def fork(self):
                """Keep each task's stage fields independent."""
                return copy.copy(self)

            def ask(self, prompt, schema, images=()):
                """Require two extraction calls to overlap."""
                if schema == EXTRACTION:
                    barrier.wait(timeout=5)
                return super().ask(prompt, schema, images)

        from reconciliation.review_settings import DEFAULTS
        with patch("reconciliation.vision_workflow.active_config", return_value={**DEFAULTS, "max_parallel": 2}):
            workflow.run(self.work, index, state, ParallelReviewer())
        self.assertEqual(len(state["units"]), 2)
        self.assertEqual(len(workflow.load(self.work)[1]["units"]), 2)

    def test_parallel_failure_keeps_other_finished_work(self):
        """Drain active jobs and save successes before reporting a budget stop."""
        barrier = threading.Barrier(2)
        saved = []

        def stopped():
            """Simulate one worker reaching the shared call budget."""
            barrier.wait(timeout=5)
            raise BudgetReached("limit")

        def finished():
            """Return another worker's completed review."""
            barrier.wait(timeout=5)
            return "finished"

        with self.assertRaises(BudgetReached):
            workflow.run_jobs([stopped, finished], saved.append, 2)
        self.assertEqual(saved, ["finished"])

    def test_omitted_pairs_cannot_pass(self):
        index, state = self.prepared()
        class OmitReviewer(FakeReviewer):
            reads = 0

            def ask(self, prompt, schema, images=()):
                if schema == EXTRACTION:
                    self.reads += 1
                    result = super().ask(prompt, schema, images)
                    result["invoice_numbers"] = [f"OTHER-{self.reads}"]
                    result["references"] = [f"OTHER-{self.reads}"]
                    return result
                if schema == SCREEN:
                    return {"comparisons": []}
                return super().ask(prompt, schema, images)
        with self.assertRaisesRegex(ValueError, "omitted"):
            workflow.run(self.work, index, state, OmitReviewer())
        self.assertTrue(workflow.gate(index, state))

    def test_modified_sources_and_previews_are_rejected(self):
        index, state = self.prepared()
        (self.root / "a.png").write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "Sources changed"):
            workflow.gate(index, state)
        unit = next(iter(index["documents"].values()))["units"][0]
        Path(unit["image"]).write_bytes(b"changed preview")
        with self.assertRaisesRegex(ValueError, "image changed"):
            workflow.load(self.work)

    def test_pdf_all_pages_and_hidden_sheet_are_extracted(self):
        pdf = self.base / "pages.pdf"
        with pymupdf.open() as document:
            for n in range(2):
                document.new_page().insert_text((30, 30), f"Receipt {n + 1}")
            document.save(pdf)
        units = extract(pdf, self.base / "pdf-units", {"pdf_mode": "vision", "pictures_enabled": True,
                                                       "codex_enabled": True, "max_calls": 20})
        self.assertEqual(len(units), 2)
        self.assertTrue(all(u["image"] for u in units))
        book = Workbook()
        book.active["A1"] = "visible"
        hidden = book.create_sheet("Hidden evidence")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "MYR 100"
        xlsx = self.base / "book.xlsx"
        book.save(xlsx)
        book.close()
        units = extract(xlsx, self.base / "sheet-units")
        self.assertTrue(any("MYR 100" in u["text"] for u in units))

    def test_stage_routing_and_changed_settings_block_resume(self):
        # Exercise mixed originals, comparison routing and resume safeguards locally.
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((30, 30), "Invoice TEST-1 Total MYR 123.45")
            pdf.save(self.root / "invoice.pdf")
        book = Workbook()
        book.active["A1"] = "Invoice TEST-1 Total MYR 123.45"
        book.save(self.root / "invoice.xlsx")
        book.close()
        index, state = self.prepared()
        from reconciliation.review_settings import DEFAULTS, STAGES
        choices = {stage: {"model": stage, "reasoning": "high"} for stage in STAGES}
        config = {**DEFAULTS, "stages": choices}
        calls = []
        class RecordingReviewer(FakeReviewer):
            def ask(self, prompt, schema, images=()):
                calls.append((self.model, self.reasoning, schema, bool(images)))
                return super().ask(prompt, schema, images)
        with patch("reconciliation.vision_workflow.active_config", return_value=config):
            workflow.run(self.work, index, state, RecordingReviewer())
            self.assertEqual({c[0] for c in calls if c[2] == EXTRACTION}, {"images", "pdf", "excel"})
            self.assertTrue(all(c[0] == "comparison" for c in calls if c[2] != EXTRACTION))
            self.assertTrue(all(c[1] == "high" for c in calls))
            self.assertTrue(all(c[3] for c in calls if c[0] == "images"))
            changed = {**config, "stages": {**choices, "pdf": {"model": "different", "reasoning": "low"}}}
            with patch("reconciliation.vision_workflow.active_config", return_value=changed):
                with self.assertRaisesRegex(workflow.ReviewPending, "Model or reasoning changed"):
                    workflow.run(self.work, index, state, RecordingReviewer())


    def test_word_claim_is_prepared_and_read_by_word_stage(self):
        """Accept Word claim text and send it through the existing Word model stage."""
        from zipfile import ZipFile
        with ZipFile(self.root / "claim.docx", "w") as document:
            document.writestr("word/document.xml",
                '<w:document xmlns:w="urn:test"><w:t>Travel claim MYR 45.00</w:t></w:document>')
        index, state = self.prepared()
        digest, claim = next((key, doc) for key, doc in index["documents"].items()
                             if doc["paths"][0].endswith(".docx"))
        self.assertTrue(claim["accepted"])
        self.assertIsNone(claim["error"])
        self.assertIn("Travel claim", claim["units"][0]["text"])
        stages = []

        class WordReviewer(FakeReviewer):
            def ask(self, prompt, schema, images=()):
                """Record extraction routing without making live model calls."""
                if schema == EXTRACTION:
                    stages.append(self.stage)
                return super().ask(prompt, schema, images)

        workflow.run(self.work, index, state, WordReviewer())
        self.assertIn("word", stages)
        self.assertIn(f"{digest}:0", state["units"])


if __name__ == "__main__":
    unittest.main()
