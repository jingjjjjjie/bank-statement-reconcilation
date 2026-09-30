"""Export the final ledger using the bank statement's styled Excel layout."""

import csv
import io
from datetime import datetime
from decimal import Decimal

from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from dashboard.services.matching import final_review
from reconciliation.bank.excel import statement_workbook


def number(value):
    """Keep absent monetary evidence blank and write known amounts as Excel numbers."""
    parsed = final_review.money(value)
    return float(parsed) if parsed is not None else None


def statement_banks(banks):
    """Recover omitted layout fields in old imports only from their hash-bound master."""
    fields = ('account', 'opening_balance', 'balance')
    if all(all(bank.get(field) not in (None, '') for field in fields) for bank in banks):
        return banks
    facts_path = final_review.CACHE / 'facts.json'
    if not facts_path.exists():
        return banks
    facts = final_review.read(facts_path)
    original = {bank['transaction_id']: bank for bank in facts['banks']}
    if any(
        bank['transaction_id'] not in original
        or any(
            bank.get(field) != original[bank['transaction_id']].get(field) for field in ('date', 'direction', 'amount')
        )
        for bank in banks
    ):
        return banks
    for path in (final_review.WORKSPACE / 'duplicated/projects').glob('*/bank-output/master_statement.csv'):
        if final_review.source_hash(path) != facts.get('bank_hash'):
            continue
        with path.open(encoding='utf-8-sig', newline='') as stream:
            source = {row.get('transaction_id'): row for row in csv.DictReader(stream)}
        return [
            {
                **bank,
                **{field: bank.get(field) or source.get(bank['transaction_id'], {}).get(field, '') for field in fields},
            }
            for bank in banks
        ]
    return banks


def export_workbook(review, include_evidence_paths=True):
    """Write all transactions with OK only for current fully supported approvals."""
    data = final_review.snapshot(review)
    banks = statement_banks(data['banks'])
    if not banks:
        raise ValueError('No bank transactions to export')
    items = {item['id']: item for item in data['items']}
    rows, details = [], []
    for bank in banks:
        decision = bank['decision'] or {}
        allocations = decision.get('allocations', []) if bank['review_status'] == 'approved' else []
        descriptions = [items[a['item_id']].get('description', '') for a in allocations] if not bank['stale'] else []
        incoming = bank.get('money_in') or (bank['amount'] if bank['direction'] == 'in' else '0')
        outgoing = bank.get('money_out') or (bank['amount'] if bank['direction'] == 'out' else '0')
        rows.append(
            {
                'date': datetime.fromisoformat(bank['date']),
                'counterparty': ' / '.join(bank['parties']).upper(),
                'particular': ' | '.join(dict.fromkeys(text for text in descriptions if text)),
                'money_in': number(incoming) or None,
                'money_out': number(outgoing) or None,
                'balance': number(bank.get('balance')),
                'status': 'OK' if bank['support_status'] == 'Supporting' else '',
            }
        )
        values = [
            bank['id'],
            bank['transaction_id'],
            bank['support_status'],
            bank['review_status'],
            bank['currency'],
            number(bank['amount']),
            number(decision.get('allocated_total', '0')),
            number(decision.get('difference', bank['amount'])),
            bank['stale'],
            ' | '.join(decision.get('flags', [])),
            decision.get('reviewer', ''),
            decision.get('note', ''),
        ]
        if include_evidence_paths:
            values.append('\n'.join(dict.fromkeys(a['source_path'] for a in allocations if a.get('source_path'))))
        details.append(values)
    metadata = {**banks[0]}
    for field in ('money_in', 'money_out'):
        # Sum source amounts with decimal arithmetic before writing numeric Excel cells.
        values = [
            bank.get(field) or (bank['amount'] if bank['direction'] == field.removeprefix('money_') else '0')
            for bank in banks
        ]
        amounts = [final_review.money(value) for value in values]
        metadata['total_' + field] = float(sum(amounts, Decimal(0))) if None not in amounts else None
    defaults = review.export_defaults() if hasattr(review, 'export_defaults') else {}
    company = defaults.get('company') or data['workspace']
    book, style = statement_workbook(company, metadata, rows)
    sheet = book.active
    for row, (bank, detail) in enumerate(zip(banks, details), 6):
        text = (
            f"{bank['id']} | {bank['review_status']} | {bank['support_status']}\n"
            f"Allocated: {detail[6]} | Difference: {detail[7]}\n"
            f"Stale evidence: {bank['stale']}\n{detail[9]}\n{detail[11]}"
        )
        style.cell(sheet, row, 'status').comment = Comment(text, 'Review')
        if rows[row - 6]['balance'] is None:
            style.cell(sheet, row, 'balance').comment = Comment('Balance unavailable in saved bank evidence.', 'Source')
    headings = [
        'Bank ID',
        'Transaction ID',
        'Support status',
        'Review status',
        'Currency',
        'Bank amount',
        'Allocated amount',
        'Difference',
        'Stale evidence',
        'Flags',
        'Reviewer',
        'Notes',
    ]
    if include_evidence_paths:
        headings.append('Supporting evidence paths')
    detail_sheet = book.create_sheet('Review details')
    for row, values in enumerate([headings, *details], 1):
        for column, value in enumerate(values, 1):
            cell = detail_sheet.cell(row, column, value)
            if isinstance(value, str):
                cell.data_type = 's'
            cell.alignment = Alignment(vertical='top', wrap_text=True)
            if row == 1:
                cell.font = Font(bold=True, color='FFFFFF')
                cell.fill = PatternFill('solid', fgColor='344B3E')
            elif column in (6, 7, 8):
                cell.number_format = '#,##0.00'
    for column, heading in enumerate(headings, 1):
        detail_sheet.column_dimensions[get_column_letter(column)].width = (
            60 if heading in ('Notes', 'Supporting evidence paths') else 24
        )
    detail_sheet.freeze_panes = 'A2'
    detail_sheet.auto_filter.ref = detail_sheet.dimensions
    stream = io.BytesIO()
    book.save(stream)
    return stream.getvalue()
