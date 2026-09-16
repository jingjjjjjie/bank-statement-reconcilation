"""Read-only, structured previews for Word documents and Excel worksheets."""
from datetime import date, datetime
from math import ceil
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

WORD = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
ROWS_PER_PAGE = 40
MAX_COLUMNS = 40


def describe(path):
    """List the visual pages available for a supported Office file."""
    path = Path(path)
    if path.suffix.lower() == ".docx":
        return {"kind": "word", "pages": 1, "labels": ["Document"]}
    if path.suffix.lower() != ".xlsx":
        raise ValueError("Visual preview is available for DOCX and XLSX files")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        labels = [f"{sheet.title} · rows {start}–{min(start + ROWS_PER_PAGE - 1, max(sheet.max_row, 1))}"
                  for sheet in workbook.worksheets
                  for start in range(1, max(sheet.max_row, 1) + 1, ROWS_PER_PAGE)]
    finally:
        workbook.close()
    return {"kind": "spreadsheet", "pages": len(labels), "labels": labels}


def _paragraph(element):
    """Preserve paragraph text and basic Word emphasis."""
    style = element.find(f"{WORD}pPr/{WORD}pStyle")
    alignment = element.find(f"{WORD}pPr/{WORD}jc")
    runs = []
    for run in element.findall(f".//{WORD}r"):
        value = "".join(part.text or "" for part in run.findall(f"{WORD}t"))
        if value:
            runs.append({"text": value, "bold": run.find(f"{WORD}rPr/{WORD}b") is not None,
                         "italic": run.find(f"{WORD}rPr/{WORD}i") is not None})
    return {"type": "paragraph", "style": style.get(f"{WORD}val", "") if style is not None else "",
            "align": alignment.get(f"{WORD}val", "") if alignment is not None else "",
            "runs": runs}


def _word(path):
    """Read paragraphs and tables in document order from a DOCX package."""
    with ZipFile(path) as archive:
        body = ET.fromstring(archive.read("word/document.xml")).find(f"{WORD}body")
    blocks = []
    for element in body:
        if element.tag == f"{WORD}p":
            blocks.append(_paragraph(element))
        elif element.tag == f"{WORD}tbl":
            rows = [[" ".join("".join(run.text or "" for run in paragraph.findall(f".//{WORD}t"))
                               for paragraph in cell.findall(f".//{WORD}p"))
                     for cell in row.findall(f"{WORD}tc")]
                    for row in element.findall(f"{WORD}tr")]
            blocks.append({"type": "table", "rows": rows})
    return {"kind": "word", "blocks": blocks}


def _value(cell):
    """Return the displayed value and basic cell styling."""
    value = cell.value
    if isinstance(value, (date, datetime)):
        value = value.isoformat(sep=" ") if isinstance(value, datetime) else value.isoformat()
    color = cell.fill.fgColor if cell.fill and cell.fill.patternType == "solid" else None
    fill = color.rgb if color is not None and color.type == "rgb" else None
    return {"text": "" if value is None else str(value), "bold": bool(cell.font and cell.font.bold),
            "fill": f"#{fill[-6:]}" if isinstance(fill, str) and len(fill) == 8 else None}


def _sheet(path, page):
    """Read one bounded worksheet section and its sheet name."""
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        offset = page
        for sheet in workbook.worksheets:
            pages = ceil(max(sheet.max_row, 1) / ROWS_PER_PAGE)
            if offset < pages:
                start = offset * ROWS_PER_PAGE + 1
                end = min(start + ROWS_PER_PAGE - 1, max(sheet.max_row, 1))
                columns = min(max(sheet.max_column, 1), MAX_COLUMNS)
                rows = [[_value(cell) for cell in row]
                        for row in sheet.iter_rows(min_row=start, max_row=end, max_col=columns)]
                return {"kind": "spreadsheet", "sheet": sheet.title, "start": start,
                        "columns": columns, "truncated_columns": sheet.max_column > MAX_COLUMNS,
                        "rows": rows}
            offset -= pages
    finally:
        workbook.close()
    raise IndexError("Preview page out of range")


def page(path, number):
    """Return one structured Office preview page."""
    path = Path(path)
    if number < 0:
        raise IndexError("Preview page out of range")
    if path.suffix.lower() == ".docx":
        if number:
            raise IndexError("Preview page out of range")
        return _word(path)
    if path.suffix.lower() == ".xlsx":
        return _sheet(path, number)
    raise ValueError("Visual preview is available for DOCX and XLSX files")
