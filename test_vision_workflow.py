"""Temporary fixtures verify coverage, admin gates, and failure handling."""
import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from PIL import Image
from openpyxl import Workbook
import pymupdf

import vision_workflow as workflow
from codex_reviewer import BudgetReached, EXTRACTION, SCREEN
from document_reader import extract
from duplicate_workflow import organize


class FakeReviewer:
    model = "fixture"

    def ask(self, prompt, schema, images=()):
        # Deterministic results test orchestration, not model accuracy.
        if schema == EXTRACTION:
            return {"readable": True, "document_type": "receipt", "references": ["TEST-1"],
                    "parties": [], "dates": [], "amounts_and_currencies": ["MYR 1"],
                    "details": "fixture", "annotations_and_signatures": "none", "limitations": []}
        if schema == SCREEN:
            payload = json.loads(prompt.split("\n", 1)[1])
            return {"comparisons": [{"right_id": key, "candidate": True, "reason": "same reference"}
                                    for key in payload["right"]]}
        return {"classification": "same_document", "confidence": "high",
                "evidence": ["same test reference"], "differences": [], "limitations": []}


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        # Fixtures live outside the customer folder and clean themselves up.
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.base = Path(self.temp.name)
        self.root = self.base / "sources"
        self.root.mkdir()
        self.work = self.base / "review"
        self.manifest = self.base / "manifest.json"
        for name, color in (("a.png", "white"), ("b.png", "gray")):
            Image.new("RGB", (80, 80), color).save(self.root / name)

    def prepared(self, duplicate=False):
        if duplicate:
            (self.root / "copy.png").write_bytes((self.root / "a.png").read_bytes())
        organize(self.root, self.manifest)
        workflow.prepare(self.manifest, self.work)
        return workflow.load(self.work)

    def completed(self):
        index, state = self.prepared()
        workflow.run(self.work, index, state, FakeReviewer())
        return index, state, next(iter(state["pairs"]))

    def test_admin_required_then_keep_both_passes(self):
        index, state, pair = self.completed()
        self.assertTrue(any("admin verdict" in p for p in workflow.gate(index, state)))
        workflow.decide(self.work, index, state, pair, "keep_both", "Admin", "Separate evidence")
        self.assertEqual(workflow.gate(index, state), [])

    def test_duplicate_cleanup_requires_survivor(self):
        index, state, pair = self.completed()
        workflow.decide(self.work, index, state, pair, "keep_left", "Admin", "Confirmed same document")
        self.assertTrue(any("cleanup pending" in p for p in workflow.gate(index, state)))
        left, right = pair.split(":")
        Path(index["documents"][right]["paths"][0]).unlink()
        self.assertEqual(workflow.gate(index, state), [])
        Path(index["documents"][left]["paths"][0]).unlink()
        with self.assertRaises(ValueError):
            workflow.gate(index, state)

    def test_pass_one_blocks_model_calls(self):
        index, state = self.prepared(duplicate=True)
        self.assertTrue((self.base / "duplicated").is_dir())
        self.assertFalse((self.root / "duplicated").exists())
        self.assertEqual(len(index["documents"]), 2)
        with self.assertRaisesRegex(ValueError, "Pass-one"):
            workflow.run(self.work, index, state, FakeReviewer())
        self.assertEqual(state["units"], {})

    def test_budget_keeps_review_incomplete(self):
        index, state = self.prepared()
        class LimitedReviewer(FakeReviewer):
            def ask(self, *args, **kwargs):
                raise BudgetReached("test limit")
        with self.assertRaises(BudgetReached):
            workflow.run(self.work, index, state, LimitedReviewer())
        self.assertTrue(workflow.gate(index, state))
        self.assertTrue((self.work / "report.md").exists())

    def test_omitted_pairs_cannot_pass(self):
        index, state = self.prepared()
        class OmitReviewer(FakeReviewer):
            def ask(self, prompt, schema, images=()):
                if schema == SCREEN:
                    return {"comparisons": []}
                return super().ask(prompt, schema, images)
        with self.assertRaisesRegex(ValueError, "omitted"):
            workflow.run(self.work, index, state, OmitReviewer())
        self.assertTrue(workflow.gate(index, state))

    def test_modified_sources_and_previews_are_rejected(self):
        index, state = self.prepared()
        (self.root / "a.png").write_bytes(b"changed")
        with self.assertRaisesRegex(ValueError, "Sources changed"):
            workflow.gate(index, state)
        unit = next(iter(index["documents"].values()))["units"][0]
        Path(unit["image"]).write_bytes(b"changed preview")
        with self.assertRaisesRegex(ValueError, "image changed"):
            workflow.load(self.work)

    def test_pdf_all_pages_and_hidden_sheet_are_extracted(self):
        pdf = self.base / "pages.pdf"
        with pymupdf.open() as document:
            for n in range(2):
                document.new_page().insert_text((30, 30), f"Receipt {n + 1}")
            document.save(pdf)
        units = extract(pdf, self.base / "pdf-units", {"pdf_mode": "vision", "pictures_enabled": True,
                                                       "codex_enabled": True, "max_calls": 20})
        self.assertEqual(len(units), 2)
        self.assertTrue(all(u["image"] for u in units))
        book = Workbook()
        book.active["A1"] = "visible"
        hidden = book.create_sheet("Hidden evidence")
        hidden.sheet_state = "hidden"
        hidden["A1"] = "MYR 100"
        xlsx = self.base / "book.xlsx"
        book.save(xlsx)
        book.close()
        units = extract(xlsx, self.base / "sheet-units")
        self.assertTrue(any("MYR 100" in u["text"] for u in units))

    def test_stage_routing_and_changed_settings_block_resume(self):
        # Exercise mixed originals, comparison routing and resume safeguards locally.
        with pymupdf.open() as pdf:
            pdf.new_page().insert_text((30, 30), "Invoice TEST-1 Total MYR 123.45")
            pdf.save(self.root / "invoice.pdf")
        book = Workbook()
        book.active["A1"] = "Invoice TEST-1 Total MYR 123.45"
        book.save(self.root / "invoice.xlsx")
        book.close()
        index, state = self.prepared()
        from review_settings import DEFAULTS, STAGES
        choices = {stage: {"model": stage, "reasoning": "high"} for stage in STAGES}
        config = {**DEFAULTS, "stages": choices}
        calls = []
        class RecordingReviewer(FakeReviewer):
            def ask(self, prompt, schema, images=()):
                calls.append((self.model, self.reasoning, schema, bool(images)))
                return super().ask(prompt, schema, images)
        with patch("vision_workflow.active_config", return_value=config):
            workflow.run(self.work, index, state, RecordingReviewer())
            self.assertEqual({c[0] for c in calls if c[2] == EXTRACTION}, {"images", "pdf", "excel"})
            self.assertTrue(all(c[0] == "comparison" for c in calls if c[2] != EXTRACTION))
            self.assertTrue(all(c[1] == "high" for c in calls))
            self.assertTrue(all(c[3] for c in calls if c[0] == "images"))
            changed = {**config, "stages": {**choices, "pdf": {"model": "different", "reasoning": "low"}}}
            with patch("vision_workflow.active_config", return_value=changed):
                with self.assertRaisesRegex(workflow.ReviewPending, "Model or reasoning changed"):
                    workflow.run(self.work, index, state, RecordingReviewer())


if __name__ == "__main__":
    unittest.main()
